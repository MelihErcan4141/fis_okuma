from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from validator.receipt_validator import ValidationStatus


SAFE_STATUSES = {
    ValidationStatus.VERIFIED,
    ValidationStatus.VERIFIED_WITH_OCR,
    ValidationStatus.VERIFIED_WITH_AI,
    ValidationStatus.NOT_PRESENT,
}


class ExcelExporter:
    def export(self, results, output_path):
        output_path = Path(output_path)

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Fisler"

        headers = [
            "Dosya",
            "Firma",
            "Fiş No",
            "Tarih",
            "Saat",
            "KDV",
            "Toplam",
            "Durum",
            "Sorunlu Alanlar",
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for item in results:
            file_path = item["file"]

            # İşleme tamamen hata verdiyse yine Excel'e yaz.
            if not item["success"]:
                sheet.append([
                    file_path,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "İŞLEME HATASI",
                    item.get("error", "")
                ])
                continue

            validation_result = item["result"]

            values = {}
            problematic_fields = []

            for field in validation_result.fields:
                if field.status in SAFE_STATUSES:
                    values[field.field_name] = (
                        field.final_value
                    )
                else:
                    values[field.field_name] = None
                    problematic_fields.append(
                        field.field_name
                    )

            if validation_result.requires_review:
                status = "KONTROL GEREKİYOR"
            else:
                status = "TAMAM"

            sheet.append([
                file_path,
                values.get("Firma"),
                values.get("Fiş No"),
                values.get("Tarih"),
                values.get("Saat"),
                values.get("KDV"),
                values.get("Toplam"),
                status,
                ", ".join(problematic_fields),
            ])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        widths = {
            "A": 35,
            "B": 40,
            "C": 15,
            "D": 15,
            "E": 12,
            "F": 12,
            "G": 15,
            "H": 22,
            "I": 35,
        }

        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        workbook.save(output_path)